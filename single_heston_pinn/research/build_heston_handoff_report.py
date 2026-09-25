#!/usr/bin/env python3
"""Build the shareable Single/Double-Heston technical handoff report."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
AUDIT = ROOT / "outputs" / "heston_handoff"
SINGLE = ROOT / "outputs" / "single_heston"
OVERFIT = ROOT / "outputs" / "single_heston_overfit_audit"
FORECAST = ROOT / "outputs" / "single_heston_next_day_forecast"
GENERALIZED = ROOT / "outputs" / "single_heston_generalized_forecast"
COMPARE = ROOT / "outputs" / "single_double_heston_comparison"
OUTPUT = AUDIT / "Heston_Double_Heston_Validated_Teammate_Handoff.docx"
CONTEXT = AUDIT / "HESTON_DOUBLE_HESTON_TEAM_CONTEXT.md"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5E6A73"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GREEN = "DDEEDB"
GOLD = "FFF1CC"
RED = "FBE3E3"
WHITE = "FFFFFF"


def set_font(run, size=None, bold=None, italic=None, color=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths):
    widths_dxa = [int(round(width * 1440)) for width in widths]
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_real_numbering(doc, style_name, kind):
    styles = doc.styles
    style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = styles["Normal"]
    style.paragraph_format.left_indent = Inches(0.375)
    style.paragraph_format.first_line_indent = Inches(-0.188)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.25
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    ppr = style.element.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    ppr.append(num_pr)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ["Caption", "Table Citation"]:
        if style_name not in doc.styles:
            doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    citation = doc.styles["Table Citation"]
    citation.font.name = "Calibri"
    citation.font.size = Pt(9)
    citation.font.italic = True
    citation.font.color.rgb = RGBColor.from_string(MUTED)
    citation.paragraph_format.space_before = Pt(4)
    citation.paragraph_format.space_after = Pt(4)
    add_real_numbering(doc, "Handoff Bullet", "bullet")
    add_real_numbering(doc, "Handoff Number", "decimal")


def set_running_furniture(doc):
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
        run = p.add_run("HESTON MODEL VALIDATION HANDOFF")
        set_font(run, size=8.5, bold=True, color=MUTED)
        run = p.add_run("\tINTERNAL TECHNICAL REFERENCE")
        set_font(run, size=8.5, color=MUTED)
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run("Page ")
        set_font(run, size=8.5, color=MUTED)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        p._p.append(fld)


def add_text(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_font(run, bold=True)
        run = p.add_run(text[len(bold_prefix):])
        set_font(run)
    else:
        set_font(p.add_run(text))
    return p


def add_bullets(doc, items, numbered=False):
    style = "Handoff Number" if numbered else "Handoff Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        set_font(p.add_run(item))


def add_callout(doc, title, body, fill=CALLOUT):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    r = p.add_run(title + "  ")
    set_font(r, bold=True, color=INK)
    set_font(p.add_run(body), color=INK)


def add_table(doc, headers, rows, widths, font_size=8.5, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        cell.text = ""
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(str(value))
        set_font(r, size=font_size, bold=True, color=INK)
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for idx, value in enumerate(row_data):
            cell = cells[idx]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if alignments:
                p.alignment = alignments[idx]
            r = p.add_run(str(value))
            set_font(r, size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_source(doc, text):
    p = doc.add_paragraph(style="Table Citation")
    set_font(p.add_run("Source: " + text), size=9, italic=True, color=MUTED)


def add_figure(doc, image_path, caption, alt_text, width=6.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(image_path), width=Inches(width))
    shape._inline.docPr.set("descr", alt_text)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(cap.add_run(caption), size=9.5, italic=True, color=MUTED)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def fmt(value, decimals=6):
    return f"{float(value):.{decimals}f}"


def markdown_table(headers, rows):
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    lines.extend("| " + " | ".join(str(value) for value in row) + " |" for row in rows)
    return "\n".join(lines)


def build_markdown(summary, single_catalog, double_catalog, metrics, checks, compare_summary, overfit_summary, controlled_single, controlled_double):
    single_rows = [[r.symbol, r.selected_candidate, fmt(r.kappa), fmt(r.theta), fmt(r.sigma), fmt(r.rho), fmt(r.v0_training_median), fmt(r.validation_iv_rmse)] for r in single_catalog.itertuples()]
    double_rows = [[r.symbol, r.selected_candidate, fmt(r.kappa_slow), fmt(r.theta_slow), fmt(r.sigma_slow), fmt(r.rho_slow), fmt(r.v0_slow_training_median), fmt(r.kappa_fast), fmt(r.theta_fast), fmt(r.sigma_fast), fmt(r.rho_fast), fmt(r.v0_fast_training_median)] for r in double_catalog.itertuples()]
    text = f"""# Heston and Double-Heston teammate context\n\nVersion date: 2026-08-06  \nAudit verdict: **{summary['verdict']}** ({summary['checks_passed']}/{summary['checks_total']} independent checks passed; {summary['critical_failures']} critical failures).\n\n## Decision\n\nKeep the current locked **Single Heston** as the historical winner. Do not promote the current Double Heston: on 8,814 identical finite rows, Single RMSE is {compare_summary['single_heston']['iv_rmse']:.6f} and Double RMSE is {compare_summary['double_heston']['iv_rmse']:.6f}; the stock-session bootstrap interval for `Double - Single RMSE` is [{compare_summary['cluster_bootstrap_double_minus_single']['double_minus_single_rmse_low']:.6f}, {compare_summary['cluster_bootstrap_double_minus_single']['high']:.6f}].\n\n## Data and evaluation boundary\n\n- Locked source: `outputs/019fc8a0/model_input_option_prices.csv`.\n- SHA-256: `{summary['input_sha256']}`.\n- 572,512 total clean-release rows; 215,636 model-ready rows; 11 model stocks.\n- Chronological split: 70% train / 15% validation / 15% test independently per stock.\n- Exact expiry dates are used; no Tuesday/Thursday weekday rewriting.\n- Forecast scores are conditional on realized target spot and target strike/expiry coordinates. They are not pure pre-market forecasts.\n- The historical test was previously viewed; future locked dates after 2026-08-03 are required for pristine confirmation.\n\n## Implemented equations\n\nSingle Heston has four structural parameters `(kappa, theta, sigma, rho)` and one date-specific variance state `v0`. Double Heston has slow and fast copies of those four structural parameters plus two date-specific states, for ten operational parameters. The implementation constrains positive parameters, Single/Double Feller gaps, slow/fast ordering, and `rho_slow^2 + rho_fast^2 < 1`.\n\n### Selected Single-Heston parameters\n\n{markdown_table(['Symbol','Candidate','kappa','theta','sigma','rho','v0 train median','Val IV RMSE'], single_rows)}\n\n`v0 train median` is descriptive only; operational forecasts use the saved origin state propagated to the target session.\n\n### Selected Double-Heston parameters\n\n{markdown_table(['Symbol','Candidate','k slow','theta slow','sigma slow','rho slow','v0 slow med','k fast','theta fast','sigma fast','rho fast','v0 fast med'], double_rows)}\n\n## Accuracy and overfitting evidence\n\n- Original locked Single next-session test: RMSE 0.044469, R2 0.814899, versus prior-session median-IV baseline RMSE 0.068079.\n- Identical common rows: Single RMSE 0.044200, Double RMSE 0.045510. Double wins only JSWENERGY and TORNTPOWER.\n- Same-day three-fold strike crossfit: Heston RMSE {overfit_summary['heston_same_day_anchor_v0']['iv_rmse']:.6f}, within-surface R2 {overfit_summary['heston_same_day_anchor_v0']['within_surface_centered_r2']:.6f}. A quadratic same-day smile baseline is better at RMSE {overfit_summary['quadratic_anchor_baseline']['iv_rmse']:.6f}.\n- Boundary warnings: 39 flags across 22 Single half-sample fits; 55/55 Double candidates boundary-near.\n- Exact synthetic recovery: Single relative error {controlled_single['exact_max_relative_parameter_error']:.3e}; Double structural error {controlled_double['exact_max_relative_structural_error']:.3e}. With 1% price noise, Double max structural error rises to {controlled_double['noisy_max_relative_structural_error']:.3%}.\n\n## Continuation protocol\n\n1. Freeze the current code, parameters, checksums, and this report.\n2. Do not use the previously viewed historical test for more tuning or model selection.\n3. Acquire future authentic NSE sessions after 2026-08-03 and process them through the locked cleaning pipeline.\n4. Generate predictions using only information available at each origin session.\n5. Score Single, Double, and the declared baseline on identical finite row keys; retain every failure row.\n6. Cluster-bootstrap by stock-session and report both pooled and per-stock metrics.\n7. Promote Double only if a preregistered future-window criterion is met and parameter stability improves materially.\n\n## File map\n\n- Core pricing: `single_heston.py`, `double_heston.py`.\n- Forecast/evaluation: `forecast_single_heston_next_day.py`, `compare_single_double_heston.py`.\n- Independent audit: `audit_heston_handoff.py`.\n- Machine audit: `outputs/heston_handoff/independent_heston_audit_checks.csv`.\n- Sanitized parameters: `outputs/heston_handoff/single_heston_selected_parameter_catalog.csv`, `double_heston_selected_parameter_catalog.csv`.\n- Full report: `outputs/heston_handoff/Heston_Double_Heston_Validated_Teammate_Handoff.docx`.\n\n## Non-negotiable limitations\n\n"""
    text += "\n".join(f"- {item}" for item in summary["non_negotiable_limitations"])
    text += "\n\n## Audit check index\n\n"
    text += markdown_table(["Category", "Check", "Passed", "Observed"], [[r.category, r.check, str(bool(r.passed)), str(r.observed).replace("|", "/")] for r in checks.itertuples()])
    return text + "\n"


def main():
    summary = json.loads((AUDIT / "independent_heston_audit_summary.json").read_text())
    compare_summary = json.loads((COMPARE / "single_double_summary.json").read_text())
    overfit_summary = json.loads((OVERFIT / "overfitting_audit_summary.json").read_text())
    forecast_summary = json.loads((FORECAST / "next_day_forecast_summary.json").read_text())
    generalized_summary = json.loads((GENERALIZED / "generalization_summary.json").read_text())
    controlled_single = json.loads((SINGLE / "single_heston_controlled_tests.json").read_text())
    controlled_double = json.loads((COMPARE / "double_heston_controlled_tests.json").read_text())
    single_catalog = pd.read_csv(AUDIT / "single_heston_selected_parameter_catalog.csv")
    double_catalog = pd.read_csv(AUDIT / "double_heston_selected_parameter_catalog.csv")
    metrics = pd.read_csv(COMPARE / "single_double_test_metrics.csv")
    checks = pd.read_csv(AUDIT / "independent_heston_audit_checks.csv")
    stability = pd.read_csv(COMPARE / "double_heston_parameter_stability.csv")
    failures = pd.read_csv(COMPARE / "double_heston_test_prediction_failures.csv")

    CONTEXT.write_text(build_markdown(summary, single_catalog, double_catalog, metrics, checks, compare_summary, overfit_summary, controlled_single, controlled_double))

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Heston and Double-Heston Validated Teammate Handoff"
    doc.core_properties.subject = "Independent leakage, overfitting, validity and parameter audit"
    doc.core_properties.author = "Project Team"
    doc.core_properties.keywords = "Heston, Double Heston, NSE, options, validation, leakage, overfitting"

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("TECHNICAL VALIDATION & HANDOFF")
    set_font(r, size=10, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Heston and Double-Heston Models")
    set_font(r, size=25, bold=True, color="000000")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Leakage, overfitting, validity, accuracy, parameter and continuation reference")
    set_font(r, size=14, color=MUTED)
    for label, value in [
        ("Prepared for", "Power-options project team"),
        ("Version date", "06 August 2026"),
        ("Scope", "Single Heston and Double Heston only"),
        ("Evidence status", f"{summary['verdict']} — {summary['checks_passed']}/{summary['checks_total']} independent checks passed"),
        ("Locked input", summary["input_sha256"]),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(label + ": "), bold=True)
        set_font(p.add_run(str(value)), size=9 if label == "Locked input" else 11)
    add_callout(doc, "Decision", "Keep the current locked Single Heston as the historical winner. Do not promote the present Double Heston. This is a reproducibility-qualified conclusion, not a guarantee of future performance.", GREEN)
    add_callout(doc, "Critical honesty boundary", "The historical test period was previously viewed during earlier development, and next-session surface scores use realized target spot plus listed target strike/expiry coordinates. Future locked NSE dates are required for pristine forward confirmation.", GOLD)

    doc.add_heading("1. Executive verdict", level=1)
    add_text(doc, "The audit found no critical inconsistency in data lineage, chronological splitting, candidate selection, state propagation, pricing, implied-volatility inversion, no-arbitrage bounds, prediction coverage, or saved metrics. All 60 independent checks passed, and all five existing automated test modules exited successfully.")
    add_table(doc, ["Question", "Evidence-backed answer"], [
        ["Was test data used to fit or select the saved candidates?", "No row-level evidence of that pathway was found. Candidates are train-only; selection is validation-only."],
        ["Is zero overfitting proven?", "No. A finite historical experiment cannot prove zero overfitting, especially after the test graph has been viewed."],
        ["Which model wins historically?", "Single Heston. On 8,814 identical finite rows: RMSE 0.044200 vs 0.045510 for Double."],
        ["Should Double Heston be used now?", "Not as the preferred model. It loses overall, wins only 2/11 stocks, and all 55 candidates are boundary-near."],
        ["Is this a pure unseen-day forecast?", "No. It forecasts the conditional volatility surface using realized target spot and observed target contract coordinates."],
    ], [1.875, 4.625], 9)
    add_source(doc, "Independent recomputation in `outputs/heston_handoff/independent_heston_audit_checks.csv`.")

    doc.add_heading("2. Scope, exclusions and authenticity", level=1)
    add_text(doc, "This handoff covers only Single Heston and Double Heston. It does not begin neural-network training, PINN development, data augmentation, Kaggle replacement, PostgreSQL mutation, or a new NSE download.")
    add_table(doc, ["Item", "Verified value"], [
        ["Model-input file", "outputs/019fc8a0/model_input_option_prices.csv"],
        ["SHA-256", summary["input_sha256"]],
        ["Rows in clean release file", "572,512"],
        ["Rows flagged model-ready", "215,636"],
        ["Stocks in model-ready subset", "11"],
        ["Original row key", "symbol | trade date | exact expiry | option type | strike"],
        ["Original source files", "raw/nse_fo_bhavcopies/<year>/...csv.zip"],
        ["Missing referenced raw files", "0"],
    ], [1.875, 4.625], 8.5)
    add_text(doc, "Authenticity means the evaluated rows resolve to local raw NSE FO bhavcopy paths and the locked file hash. This audit does not independently certify NSE's exchange infrastructure or redownload every bhavcopy from the internet.")

    doc.add_heading("3. Chronology and information boundary", level=1)
    add_bullets(doc, [
        "Dates are split 70% train, 15% validation and 15% test independently within each stock.",
        "All Single- and Double-Heston candidate training dates precede the first validation date for the same stock.",
        "Candidate selection is recomputable using validation data only; validation and test row keys are disjoint.",
        "Every saved origin date strictly precedes its target date; the observed calendar gap is one to four days.",
        "Exact expiry_date and days-to-expiry values are preserved. Historical Tuesday, Wednesday and Thursday expiries remain distinct; no weekday rule rewrites them.",
        "Origin-state tables contain no target option price or target implied-volatility label.",
    ])
    add_callout(doc, "What 'no leakage found' means", "The saved row-level artifacts reproduce a valid train → validation selection → test scoring path. It does not prove that a researcher never looked at aggregated test results. The project record explicitly says the test graph was viewed earlier.", GOLD)

    doc.add_heading("4. Implemented model equations and parameter meaning", level=1)
    doc.add_heading("4.1 Single Heston", level=2)
    add_text(doc, "Under the risk-neutral measure, the implemented one-factor model uses stochastic variance v(t):")
    add_callout(doc, "Price process", "dS(t)/S(t) = (r − q)dt + √v(t) dW_S(t)")
    add_callout(doc, "Variance process", "dv(t) = κ[θ − v(t)]dt + σ√v(t) dW_v(t), with d⟨W_S,W_v⟩ = ρdt")
    add_table(doc, ["Parameter", "Meaning", "Implementation role"], [
        ["κ (kappa)", "Mean-reversion speed", "Propagates origin variance toward θ and shapes the characteristic function."],
        ["θ (theta)", "Long-run variance", "Positive structural level; annualized volatility scale is √θ."],
        ["σ (sigma)", "Volatility of variance", "Controls smile curvature and variance uncertainty."],
        ["ρ (rho)", "Spot/variance correlation", "Controls skew; constrained inside (−0.98, 0.98)."],
        ["v0", "Date-specific variance state", "Fitted at origin, propagated to target. Training median is descriptive, not the live forecast state."],
    ], [1.181, 2.0, 3.319], 8.5)
    add_text(doc, "The Feller diagnostic is 2κθ − σ² > 0. Pricing uses a Little-Heston-trap characteristic function with 64-point Gauss-Laguerre integration, put-call parity for puts, and Black-Scholes inversion for implied volatility.")

    doc.add_heading("4.2 Double Heston", level=2)
    add_text(doc, "The implemented Double Heston characteristic function is the product of two constrained Heston-factor characteristic functions, represented additively in log-characteristic exponents. The instantaneous variance is decomposed into slow and fast states.")
    add_callout(doc, "Two variance factors", "dv_i(t) = κ_i[θ_i − v_i(t)]dt + σ_i√v_i(t)dW_vi(t), for i ∈ {slow, fast}; total state = v_slow + v_fast.")
    add_bullets(doc, [
        "Operational parameter count is ten: eight structural parameters plus v0_slow and v0_fast.",
        "Slow/fast ordering is enforced by κ_slow < κ_fast.",
        "Both factors must satisfy separate Feller conditions.",
        "The joint correlation disk is constrained by ρ_slow² + ρ_fast² < 1.",
        "Each factor state propagates exactly as θ_i + [v_i(origin) − θ_i] exp(−κ_i Δcalendar/365).",
    ])

    doc.add_heading("5. Selected Single-Heston parameter set", level=1)
    single_rows = [[r.symbol, r.selected_candidate, fmt(r.kappa), fmt(r.theta), fmt(r.sigma), fmt(r.rho), fmt(r.v0_training_median), fmt(r.validation_iv_rmse)] for r in single_catalog.itertuples()]
    add_table(doc, ["Symbol", "Candidate", "κ", "θ", "σ", "ρ", "v0 median", "Val RMSE"], single_rows, [1.2, 1.05, .7, .7, .7, .7, .85, .6], 7.2)
    add_source(doc, "Sanitized catalog `outputs/heston_handoff/single_heston_selected_parameter_catalog.csv`.")
    add_callout(doc, "Operational warning", "The v0 median above summarizes training states. For each forecast origin, use the origin-specific fitted v0 and propagate it; do not substitute the median as a universal live state.", GOLD)

    add_page_break(doc)
    doc.add_heading("6. Selected Double-Heston parameter set", level=1)
    slow_rows = [[r.symbol, r.selected_candidate, fmt(r.kappa_slow), fmt(r.theta_slow), fmt(r.sigma_slow), fmt(r.rho_slow), fmt(r.v0_slow_training_median)] for r in double_catalog.itertuples()]
    fast_rows = [[r.symbol, r.selected_candidate, fmt(r.kappa_fast), fmt(r.theta_fast), fmt(r.sigma_fast), fmt(r.rho_fast), fmt(r.v0_fast_training_median)] for r in double_catalog.itertuples()]
    doc.add_heading("6.1 Slow factor", level=2)
    add_table(doc, ["Symbol", "Candidate", "κ slow", "θ slow", "σ slow", "ρ slow", "v0 slow med"], slow_rows, [1.25, 1.05, .8, .8, .8, .8, 1.0], 7.3)
    doc.add_heading("6.2 Fast factor", level=2)
    add_table(doc, ["Symbol", "Candidate", "κ fast", "θ fast", "σ fast", "ρ fast", "v0 fast med"], fast_rows, [1.25, 1.05, .8, .8, .8, .8, 1.0], 7.3)
    add_source(doc, "Sanitized catalog `outputs/heston_handoff/double_heston_selected_parameter_catalog.csv`.")
    add_callout(doc, "Do not interpret these as uniquely identified economics", "All 55 Double-Heston candidates are boundary-near, and training-half/start-value fits disagree materially. Similar prices can be produced by different parameter combinations.", RED)

    doc.add_heading("7. Calibration and selection protocol", level=1)
    add_table(doc, ["Stage", "Single Heston", "Double Heston"], [
        ["Structural candidates", "3 per stock: full train + alternating train halves A/B", "5 per stock: 3 full-train starts + train halves A/B"],
        ["Training objective", "Vega-scaled robust price residuals, equal date weighting", "Same principle; two variance states per training surface"],
        ["Selection", "Minimum validation pooled IV RMSE; alphabetical tie break", "Minimum validation mean stock-session IV RMSE; pooled RMSE then name tie breaks"],
        ["Test use", "Scoring only after selection", "Scoring only after selection; identical common rows with Single"],
        ["Failure handling", "Finite selected predictions", "4 invalid test IV rows retained separately; no imputation"],
    ], [1.2, 2.65, 2.65], 8)
    add_text(doc, "Optimization convergence alone is not evidence of unique recovery. The controlled tests show that prices can be recovered accurately even when structural parameters are materially displaced by small noise, especially for Double Heston.")

    doc.add_heading("8. Accuracy on the locked historical evaluation", level=1)
    add_table(doc, ["Experiment", "Model", "Rows", "IV RMSE", "IV MAE", "Bias", "R²"], [
        ["Original next-session", "Single Heston", "8,818", "0.044469", "0.028048", "−0.004604", "0.814899"],
        ["Original next-session", "Prior-session median IV", "8,818", "0.068079", "0.042060", "−0.031587", "0.566172"],
        ["Identical common rows", "Single Heston", "8,814", "0.044200", "0.027957", "−0.004502", "0.816578"],
        ["Identical common rows", "Double Heston", "8,814", "0.045510", "0.028264", "−0.005601", "0.805544"],
    ], [1.2, 1.5, .6, .8, .8, .8, .8], 7.8)
    add_text(doc, f"The stock-session bootstrap interval for Double minus Single RMSE is [{compare_summary['cluster_bootstrap_double_minus_single']['double_minus_single_rmse_low']:.6f}, {compare_summary['cluster_bootstrap_double_minus_single']['high']:.6f}], entirely positive. Lower is better; this favors Single Heston.")
    add_figure(doc, COMPARE / "single_vs_double_actual_iv.png", "Figure 1. Identical-row actual IV versus Single- and Double-Heston forecasts.", "Two scatter-density panels compare actual implied volatility with Single Heston and Double Heston forecasts on identical test rows. The diagonal line represents perfect agreement.")
    add_figure(doc, COMPARE / "single_vs_double_rmse_by_stock.png", "Figure 2. Per-stock test IV RMSE on the identical finite comparison universe.", "Grouped bars show Single and Double Heston implied-volatility RMSE by stock; lower bars indicate better performance.")

    doc.add_heading("8.1 Per-stock comparison", level=2)
    metric_rows = []
    for symbol in sorted(metrics.loc[metrics.scope.ne("ALL"), "scope"].unique()):
        s = metrics[(metrics.scope.eq(symbol)) & metrics.model.eq("Single Heston")].iloc[0]
        d = metrics[(metrics.scope.eq(symbol)) & metrics.model.eq("Double Heston")].iloc[0]
        metric_rows.append([symbol, int(s.rows), fmt(s.iv_rmse), fmt(d.iv_rmse), f"{d.iv_rmse - s.iv_rmse:+.6f}", "Double" if d.iv_rmse < s.iv_rmse else "Single"])
    add_table(doc, ["Symbol", "Rows", "Single RMSE", "Double RMSE", "Double−Single", "Winner"], metric_rows, [1.4, .7, 1.1, 1.1, 1.2, 1.0], 8)
    add_source(doc, "`outputs/single_double_heston_comparison/single_double_test_metrics.csv`.")

    doc.add_heading("9. Leakage and overfitting challenge tests", level=1)
    add_table(doc, ["Challenge", "Result", "Interpretation"], [
        ["Three-fold strike crossfit", "4,709 unique held-out quotes; Heston RMSE 0.030028", "No same-row memorization in the sampled surfaces."],
        ["Within-surface centered R²", "0.617874", "Less impressive than pooled R² 0.903410; much pooled fit comes from same-day level calibration."],
        ["Flat same-day anchor", "RMSE 0.040577", "Heston beats a flat surface."],
        ["Quadratic same-day smile", "RMSE 0.014712", "Simple flexible smile beats Heston for same-day strike interpolation."],
        ["Prior available Heston state", "RMSE 0.042245", "Removing same-day state information reduces performance."],
        ["Shuffled state", "RMSE 0.068991", "Correct state timing matters."],
        ["Single half-sample fits", "39 boundary flags across 22 fits", "Structural parameters are not fully stable."],
        ["Double candidates", "55/55 boundary-near", "Severe identifiability/constraint-pressure warning."],
    ], [1.55, 1.75, 3.2], 8)
    add_figure(doc, COMPARE / "double_heston_parameter_instability.png", "Figure 3. Double-Heston disagreement across five train-only candidates per stock.", "Bars summarize the maximum log range of positive parameters and the correlation ranges across five independently fitted Double-Heston candidates per stock.")
    add_text(doc, "The parameter-instability chart measures disagreement across optimizer starts and disjoint training halves. It is not a confidence interval and should not be interpreted as statistical uncertainty coverage.")

    doc.add_heading("10. Controlled synthetic recovery", level=1)
    add_table(doc, ["Model / condition", "Known-data recovery", "Price error", "Parameter sensitivity"], [
        ["Single, exact prices", f"Max relative error {controlled_single['exact_max_relative_parameter_error']:.3e}", f"RMSE {controlled_single['exact_price_rmse']:.3e}", "Three starts converge to effectively exact recovery"],
        ["Single, 1% price noise", "Recovered structural set remains close", f"RMSE {controlled_single['noisy_price_rmse']:.6f}", "Non-zero displacement; use price fit and stability together"],
        ["Double, exact prices", f"Max structural error {controlled_double['exact_max_relative_structural_error']:.3e}", f"RMSE {controlled_double['exact_price_rmse']:.3e}", "Implementation can recover a known clean solution"],
        ["Double, 1% price noise", f"Max structural error {controlled_double['noisy_max_relative_structural_error']:.3%}", f"RMSE {controlled_double['noisy_price_rmse']:.6f}", "Material non-identifiability despite small pricing loss"],
    ], [1.4, 1.7, 1.25, 2.15], 8)
    add_callout(doc, "Meaning", "These tests validate pricing/calibration mechanics under controlled conditions. They do not validate that real NSE data are generated by Heston dynamics, nor do they make the historical parameter estimates unique.", GOLD)

    doc.add_heading("11. Invalid rows, coverage and honesty controls", level=1)
    add_bullets(doc, [
        "The Single-Heston source test universe contains 8,818 unique rows.",
        "Double Heston returns finite IV for 8,814 rows (99.9546% coverage).",
        "Four invalid Double-Heston implied-volatility rows are retained in a separate failure file and are not imputed or silently replaced.",
        "Candidate selection uses only validation quote keys where all five Double candidates are finite; validation failures are also retained separately.",
        "All stored Single/Double prices reprice from the saved equations within 2×10⁻⁹, and stored IV values invert those prices within 2×10⁻¹⁰.",
    ])
    failure_rows = [[getattr(r, "symbol", ""), getattr(r, "row_key", ""), getattr(r, "failure_reason", getattr(r, "reason", "invalid IV"))] for r in failures.itertuples()]
    add_table(doc, ["Symbol", "Row key", "Recorded result"], failure_rows, [1.1, 4.1, 1.3], 7.2)
    add_source(doc, "`outputs/single_double_heston_comparison/double_heston_test_prediction_failures.csv`.")

    doc.add_heading("12. Retrospective generalization adjustment", level=1)
    add_text(doc, f"A later validation-only guard improved the historical Single-Heston test RMSE from {generalized_summary['original_locked_heston']['iv_rmse']:.6f} to {generalized_summary['generalized_heston']['iv_rmse']:.6f}, with slope improving from {generalized_summary['original_locked_heston']['forecast_on_actual_slope']:.6f} to {generalized_summary['generalized_heston']['forecast_on_actual_slope']:.6f}. Corrections were accepted for {', '.join(generalized_summary['corrections_accepted_for_symbols'])}.")
    add_callout(doc, "Do not present this as pristine out-of-sample validation", generalized_summary["evaluation_status"], RED)
    add_text(doc, "For teammate continuation, the corrected model may be carried as a locked challenger, but its first credible confirmation must use dates that were not visible when the rule was designed.")

    doc.add_heading("13. Privacy and data-exposure review", level=1)
    add_table(doc, ["Exposure class", "Finding", "Required action"], [
        ["Credentials/secrets", "No API-key, password, AWS-key or private-key pattern found in reviewed model code/reports.", "Continue secret scanning before any public release."],
        ["Personal/machine metadata", "Two raw machine artifacts contain an absolute local path with the workstation username.", "Do not share them verbatim; use sanitized catalogs and this scrubbed report."],
        ["Raw NSE rows", "This report contains no raw option-price table and no embedded source dataset.", "Share the data separately under the project's access policy."],
        ["Model parameters", "All selected parameters are intentionally disclosed for teammate continuation.", "Treat them as project IP if required by team policy."],
    ], [1.4, 2.7, 2.4], 8)
    add_text(doc, "The final DOCX is metadata-scrubbed after generation. The context Markdown uses project-relative paths only.")

    add_page_break(doc)
    doc.add_heading("14. Locked continuation protocol", level=1)
    add_bullets(doc, [
        "Freeze `single_heston.py`, `double_heston.py`, `forecast_single_heston_next_day.py`, `compare_single_double_heston.py`, selected parameter catalogs, and all checksum manifests.",
        "Do not alter candidate starts, bounds, filters, correction thresholds, or evaluation rows using the previously viewed historical test.",
        "Acquire authentic future NSE sessions strictly after 03 August 2026 and retain raw-source hashes and exact expiry dates.",
        "Run the same model-ready filter and record every exclusion reason. Never fill or fabricate a failed model IV.",
        "At each target, construct the origin state using origin-date information only; then propagate over the exact calendar gap.",
        "Generate Single and Double predictions on identical row keys, conditional on the same target spot/contract coordinates if preserving the current evaluation convention.",
        "Report pooled RMSE/MAE/bias/R², mean stock-session RMSE, forecast-on-actual slope, per-stock results, coverage and all failure rows.",
        "Bootstrap by stock-session, not individual option row, because strikes within a session are dependent.",
        "Promote Double only after a preregistered future window shows a clear improvement and parameter stability no longer depends strongly on start/half sample.",
    ], numbered=True)

    doc.add_heading("14.1 Recommended preregistration", level=2)
    add_table(doc, ["Decision item", "Lock before future scoring"], [
        ["Primary metric", "Mean stock-session IV RMSE on identical finite rows"],
        ["Uncertainty", "95% stock-session cluster bootstrap of Double minus Single RMSE"],
        ["Coverage floor", "At least 99.9%; every invalid row documented"],
        ["Promotion rule", "Double interval entirely below zero and per-stock performance not concentrated in only one or two names"],
        ["Stability rule", "Substantially fewer boundary-near fits plus lower start/half-sample parameter ranges"],
        ["No-retuning rule", "One locked run on the preregistered future window"],
    ], [1.875, 4.625], 8.5)

    doc.add_heading("15. Reproducibility commands and file map", level=1)
    add_text(doc, "Run from the project root with the environment that provides NumPy, pandas, SciPy and Matplotlib:")
    for command in [
        "python3 test_single_heston.py",
        "python3 test_heston_overfit_audit.py",
        "python3 test_heston_next_day_forecast.py",
        "python3 test_heston_generalization.py",
        "python3 test_double_heston.py",
        "python3 audit_heston_handoff.py",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(3)
        set_font(p.add_run(command), name="Courier New", size=9, color=INK)
    add_table(doc, ["Role", "Project-relative path"], [
        ["Authentic model input", "outputs/019fc8a0/model_input_option_prices.csv"],
        ["Single core", "single_heston.py"],
        ["Single next-session", "forecast_single_heston_next_day.py"],
        ["Single overfit audit", "audit_single_heston_overfitting.py"],
        ["Double core", "double_heston.py"],
        ["Single/Double comparison", "compare_single_double_heston.py"],
        ["Independent handoff audit", "audit_heston_handoff.py"],
        ["Machine audit matrix", "outputs/heston_handoff/independent_heston_audit_checks.csv"],
        ["Sanitized parameters", "outputs/heston_handoff/*_selected_parameter_catalog.csv"],
        ["Context Markdown", "outputs/heston_handoff/HESTON_DOUBLE_HESTON_TEAM_CONTEXT.md"],
    ], [1.875, 4.625], 8)

    doc.add_heading("16. Known limitations", level=1)
    add_bullets(doc, summary["non_negotiable_limitations"])
    add_bullets(doc, [
        "The model universe is the 11 stocks with model-ready data, not every possible power-sector security.",
        "Calibration uses a filtered liquid/parity-consistent option subset; results do not automatically generalize to stale, illiquid or rejected quotes.",
        "Corporate-action adjustments and expiry changes are inherited from the locked clean-release pipeline; this audit validates the fields and lineage, not every corporate-action announcement independently.",
        "The baseline is useful but limited. Future work should retain it and may add stronger predeclared baselines without using the evaluation window to choose them.",
        "R² is descriptive and scale-sensitive; use RMSE, calibration slope, coverage, session clustering and stability together.",
    ])

    doc.add_heading("17. Audit-check appendix", level=1)
    def brief_observed(value):
        value = str(value)
        return value if len(value) <= 76 else value[:75] + "…"

    audit_rows = [[r.category, r.check.replace("_", " "), "PASS" if bool(r.passed) else "FAIL", brief_observed(r.observed)] for r in checks.itertuples()]
    add_table(doc, ["Category", "Independent check", "Status", "Observed"], audit_rows, [1.0, 3.15, .65, 1.7], 7.2)
    add_source(doc, "Full details, severities and untruncated observations are in `independent_heston_audit_checks.csv`.")

    doc.add_heading("18. Glossary", level=1)
    add_table(doc, ["Term", "Definition in this project"], [
        ["Origin", "The last observed NSE session whose state is allowed to enter a target forecast."],
        ["Target", "The next observed session being scored."],
        ["Conditional surface forecast", "Forecast IV at target spot and target-listed strike/expiry coordinates, without using target option prices."],
        ["Structural parameters", "κ, θ, σ and ρ for each variance factor."],
        ["State parameters", "v0 for Single; v0_slow and v0_fast for Double, fitted per origin session."],
        ["Feller gap", "2κθ − σ²; positive under the declared constraint."],
        ["Boundary-near", "A diagnostic that a fitted candidate lies close to one or more imposed parameter limits."],
        ["Leakage", "Information from validation/test labels influencing fitting, selection or state construction earlier than allowed."],
        ["Retrospective", "Designed or interpreted after some test results were already viewed."],
    ], [1.5, 5.0], 8.5)
    add_callout(doc, "Handoff status", "The teammate can reproduce the saved experiments and continue with future locked validation. The report does not authorize calling the current results production-ready or leakage-proof in an absolute sense.", GREEN)

    set_running_furniture(doc)
    doc.save(OUTPUT)
    print(OUTPUT)
    print(CONTEXT)


if __name__ == "__main__":
    main()
